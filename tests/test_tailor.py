"""Resume tailoring safety.

The one unrecoverable failure mode is a fabricated claim reaching an employer,
so these tests target invention detection and the exact-match constraint on
edits rather than the prose quality.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from lgrow.jobs import tailor


class TestSuspiciousNumbers:
    def test_new_percentage_flagged(self):
        # The classic fabrication: a metric that was never there.
        invented = tailor.suspicious_numbers(
            "Maintained the billing service.",
            "Improved billing throughput by 40%.",
        )
        assert invented

    def test_reworded_but_same_numbers_is_clean(self):
        assert tailor.suspicious_numbers(
            "Cut p99 latency from 800ms to 200ms.",
            "Reduced p99 latency 800ms -> 200ms on the checkout path.",
        ) == []

    def test_no_numbers_anywhere_is_clean(self):
        assert tailor.suspicious_numbers(
            "Led the migration off the legacy queue.",
            "Owned the migration away from the legacy queue.",
        ) == []

    def test_dropping_a_number_is_not_suspicious(self):
        # Removing a fact is the candidate's choice; adding one is not.
        assert tailor.suspicious_numbers("Shipped 12 services.", "Shipped services.") == []

    def test_multiplier_flagged(self):
        assert tailor.suspicious_numbers("Scaled the cluster.", "Scaled the cluster 10x.")


def make_docx(path: Path, paragraphs: list[str]) -> None:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))


BASE = [
    "Jane Doe",
    "jane@example.com",
    "Experienced engineer who has worked on web services and data pipelines "
    "for six years across two companies in the payments space.",
    "Built internal tooling for deployment and monitoring.",
    "Maintained the billing service and its test suite.",
]


class TestApplyRewrites:
    def test_matching_edits_applied(self, tmp_path):
        import docx

        src = tmp_path / "base.docx"
        dest = tmp_path / "out.docx"
        make_docx(src, BASE)

        tailored = tailor.TailoredResume(
            rewrites=[
                tailor.Rewrite(
                    original="Built internal tooling for deployment and monitoring.",
                    replacement="Built internal CI/CD and observability tooling.",
                )
            ]
        )
        applied, skipped, warnings = tailor.apply_rewrites(src, dest, tailored)
        assert applied == 1
        assert skipped == []
        assert warnings == []

        texts = [p.text for p in docx.Document(str(dest)).paragraphs]
        assert "Built internal CI/CD and observability tooling." in texts
        # Untouched content must survive verbatim.
        assert "Jane Doe" in texts
        assert "jane@example.com" in texts

    def test_non_matching_edit_is_skipped_not_appended(self, tmp_path):
        """A model returning text that isn't in the resume must not add content."""
        import docx

        src = tmp_path / "base.docx"
        dest = tmp_path / "out.docx"
        make_docx(src, BASE)

        tailored = tailor.TailoredResume(
            rewrites=[
                tailor.Rewrite(
                    original="A paragraph that does not exist in the resume",
                    replacement="Led a team of 20 engineers at Google.",
                )
            ]
        )
        applied, skipped, _ = tailor.apply_rewrites(src, dest, tailored)
        assert applied == 0
        assert len(skipped) == 1

        texts = "\n".join(p.text for p in docx.Document(str(dest)).paragraphs)
        assert "Google" not in texts
        assert len(docx.Document(str(dest)).paragraphs) == len(BASE)

    def test_edit_inventing_a_metric_is_refused(self, tmp_path):
        import docx

        src = tmp_path / "base.docx"
        dest = tmp_path / "out.docx"
        make_docx(src, BASE)

        tailored = tailor.TailoredResume(
            rewrites=[
                tailor.Rewrite(
                    original="Maintained the billing service and its test suite.",
                    replacement="Cut billing incidents by 60% and raised coverage to 95%.",
                )
            ]
        )
        applied, _, warnings = tailor.apply_rewrites(src, dest, tailored)
        assert applied == 0
        assert warnings and "figures not in your resume" in warnings[0]

        texts = "\n".join(p.text for p in docx.Document(str(dest)).paragraphs)
        assert "60%" not in texts
        assert "Maintained the billing service" in texts

    def test_source_file_is_never_modified(self, tmp_path):
        import docx

        src = tmp_path / "base.docx"
        dest = tmp_path / "out.docx"
        make_docx(src, BASE)
        before = src.read_bytes()

        tailor.apply_rewrites(
            src, dest,
            tailor.TailoredResume(
                rewrites=[tailor.Rewrite(original=BASE[3], replacement="Changed.")]
            ),
        )
        assert src.read_bytes() == before
        assert "Changed." in [p.text for p in docx.Document(str(dest)).paragraphs]


class TestReadResume:
    def test_missing_file_explains_the_fix(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="master resume"):
            tailor.read_resume_paragraphs(tmp_path / "nope.docx")

    def test_reads_paragraphs_skipping_blanks(self, tmp_path):
        src = tmp_path / "r.docx"
        make_docx(src, ["One", "", "   ", "Two"])
        assert tailor.read_resume_paragraphs(src) == ["One", "Two"]
