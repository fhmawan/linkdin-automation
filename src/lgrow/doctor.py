"""Preflight diagnostics.

`lgrow doctor` answers one question: what do I have to fix before this works?
Every failure carries the exact command or file edit that resolves it, because
this is the first thing run on a fresh clone and after anything breaks.
"""

from __future__ import annotations

import json
import os
import shutil
import sqlite3
import sys
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from . import config, db, llm, paths

OK, WARN, FAIL = "ok", "warn", "fail"

_ICON = {OK: "✓", WARN: "!", FAIL: "✗"}
_COLOR = {OK: "\033[32m", WARN: "\033[33m", FAIL: "\033[31m"}
_RESET = "\033[0m"


@dataclass
class Check:
    name: str
    status: str
    detail: str = ""
    hint: str = ""


def _color_enabled() -> bool:
    return sys.stdout.isatty() and os.environ.get("NO_COLOR") is None


def render(checks: list[Check]) -> str:
    use_color = _color_enabled()
    width = max((len(c.name) for c in checks), default=0)
    lines: list[str] = []
    for c in checks:
        icon = _ICON[c.status]
        if use_color:
            icon = f"{_COLOR[c.status]}{icon}{_RESET}"
        lines.append(f"  {icon}  {c.name.ljust(width)}   {c.detail}")
        if c.hint and c.status != OK:
            for hint_line in c.hint.splitlines():
                lines.append(f"     {' ' * width}   ↳ {hint_line}")
    return "\n".join(lines)


# ─── individual checks ───────────────────────────────────────────────────────


def check_python() -> Check:
    v = sys.version_info
    detail = f"{v.major}.{v.minor}.{v.micro}"
    if (v.major, v.minor) < (3, 11):
        return Check("python", FAIL, detail, "lgrow needs Python 3.11 or newer.")
    return Check("python", OK, detail)


def configured_providers() -> set[str]:
    """Every provider named anywhere in the model chains."""
    models = config.load_ai().models
    specs = [*models.chain("quality"), *models.chain("cheap")]
    return {llm.split_spec(s)[0] for s in specs}


def check_model_chains() -> list[Check]:
    """Report each tier's chain, and flag one that cannot survive a quota wall.

    Every free tier in play caps by the day, so a chain confined to one provider
    stops entirely once that allowance is spent — and the symptom is a 429 that
    reads like a global outage rather than a config choice.
    """
    models = config.load_ai().models
    checks: list[Check] = []
    for tier in ("quality", "cheap"):
        chain = models.chain(tier)
        if not chain:
            checks.append(
                Check(f"models {tier}", FAIL, "no model configured",
                      f"Set models.{tier} in config/ai.yaml.")
            )
            continue
        providers = {llm.split_spec(s)[0] for s in chain}
        detail = f"{chain[0]} (+{len(chain) - 1} fallback)"
        if len(chain) == 1:
            checks.append(
                Check(f"models {tier}", WARN, detail,
                      f"No fallback. Add entries under models.fallbacks.{tier} in\n"
                      f"config/ai.yaml so a spent daily quota does not stop the run.")
            )
        elif len(providers) == 1:
            only = next(iter(providers))
            # Not a defect for `quality`: it is pinned to one provider on
            # purpose, because those prompts carry the resume and contact
            # details. Say what the consequence is and leave the call alone.
            note = (
                "Single provider, so the whole chain shares one daily cap. That is\n"
                "deliberate here — quality prompts carry your resume and contact\n"
                "details, so they stay with one provider you have vetted."
                if tier == "quality"
                else "Single provider, so the whole chain shares one daily cap. Adding\n"
                     "a model from another provider would survive that cap."
            )
            checks.append(
                Check(f"models {tier}", OK if tier == "quality" else WARN,
                      f"{detail}, {only} only", note)
            )
        else:
            checks.append(Check(f"models {tier}", OK,
                                f"{detail}, {len(providers)} providers"))
    return checks


def check_llm_auth() -> list[Check]:
    """A key per provider actually referenced by the model chains."""
    config.load_env()
    source = ".env" if paths.ENV_PATH.exists() else "environment"
    hints = {
        "gemini": (
            "GEMINI_API_KEY",
            "Create a key at https://aistudio.google.com/apikey.\n"
            "Personal 'Login with Google' is no longer accepted: Google retired\n"
            "the free Code Assist tier for individuals.",
        ),
        "openrouter": (
            "OPENROUTER_API_KEY",
            "Create a key at https://openrouter.ai/keys.\n"
            "Most :free endpoints also need 'free endpoints that may train on\n"
            "your inputs' enabled in your OpenRouter privacy settings.",
        ),
    }
    checks: list[Check] = []
    for provider in sorted(configured_providers()):
        var, hint = hints[provider]
        value = os.environ.get(var, "").strip()
        if value.startswith(f"{var}="):
            # Pasting the whole `NAME=value` line under an existing `NAME=`
            # leaves the name inside the value. Every presence check still
            # passes and the server returns a baffling 401.
            checks.append(
                Check(
                    f"{provider} auth", FAIL, f"{var} contains its own name",
                    f"The value in .env starts with '{var}='. You likely pasted the\n"
                    f"whole line under the existing key. It should read:\n"
                    f"    {var}=<the key itself>",
                )
            )
        elif value:
            checks.append(Check(f"{provider} auth", OK, f"via ${var} ({source})"))
        else:
            checks.append(
                Check(
                    f"{provider} auth", FAIL, f"{var} is not set",
                    f"{hint}\n"
                    f"Put it in .env rather than your shell profile — cron does not\n"
                    f"read a login profile, so a key exported there works by hand\n"
                    f"and then fails at 07:07:\n"
                    f"    {var}=...",
                )
            )
    return checks


def check_gemini_probe(tier: str = "quality") -> Check:
    ok, detail = llm.probe(tier)
    if ok:
        return Check(f"probe ({tier})", OK, detail)
    hint = "Check the API key for this tier's provider in .env."
    if "Budget" in detail:
        hint = "Raise budget.max_calls_per_day in config/ai.yaml."
    elif "model" in detail.lower() and (
        "not found" in detail.lower() or "no longer available" in detail.lower()
    ):
        hint = (
            f"The model ID for tier '{tier}' looks stale — these churn.\n"
            f"Pick a current one and update models.{tier} in config/ai.yaml:\n"
            f"  Gemini:     GET generativelanguage.googleapis.com/v1beta/models\n"
            f"  OpenRouter: GET openrouter.ai/api/v1/models"
        )
    return Check(f"probe ({tier})", FAIL, detail, hint)


def check_budget() -> Check:
    try:
        ai = config.load_ai()
        remaining = llm.remaining_budget()
    except Exception as exc:  # noqa: BLE001
        return Check("call budget", WARN, f"unreadable: {exc}")
    total = ai.budget.max_calls_per_day
    status = OK if remaining > ai.budget.reserve_calls else WARN
    return Check(
        "call budget",
        status,
        f"{remaining}/{total} calls left today (UTC)",
        "Budget nearly spent; the next batch stage will refuse to start."
        if status == WARN
        else "",
    )


def check_configs() -> list[Check]:
    checks: list[Check] = []
    loaders = (
        ("profile.yaml", config.load_profile),
        ("sources.yaml", config.load_sources),
        ("voice.yaml", config.load_voice),
        ("ai.yaml", config.load_ai),
    )
    models = {}
    for name, loader in loaders:
        try:
            models[name] = loader()
        except Exception as exc:  # noqa: BLE001
            checks.append(
                Check(f"config {name}", FAIL, str(exc).splitlines()[0][:120],
                      f"Fix the YAML in config/{name}")
            )
        else:
            checks.append(Check(f"config {name}", OK, "parsed"))

    profile = models.get("profile.yaml")
    if profile is not None:
        missing = config.unfilled(profile)
        if missing:
            checks.append(
                Check(
                    "profile filled in",
                    FAIL,
                    f"{len(missing)} placeholder(s)",
                    "Edit config/profile.yaml and replace FILL_ME at:\n"
                    + "\n".join(f"  {m}" for m in missing[:12])
                    + ("\n  ..." if len(missing) > 12 else ""),
                )
            )
        else:
            checks.append(Check("profile filled in", OK, "no placeholders"))

        if not profile.targets.roles:
            checks.append(
                Check("target roles", FAIL, "empty",
                      "Add at least one role to targets.roles in config/profile.yaml")
            )
        if not profile.targets.skills_core:
            checks.append(
                Check("core skills", WARN, "empty",
                      "targets.skills_core drives match quality — add a few.")
            )
        if not profile.git.authors:
            checks.append(
                Check(
                    "git authors",
                    WARN,
                    "empty — commit harvesting disabled",
                    "Post drafts will rely on data/journal.md only.\n"
                    "To include your commits, set git.authors in config/profile.yaml\n"
                    "to the identity YOU commit under. Find it with:\n"
                    "    git -C <your-repo> log --format='%an <%ae>' | sort -u | head",
                )
            )
        else:
            missing_repos = [p for p in profile.git.repo_paths if not Path(p).expanduser().is_dir()]
            if missing_repos:
                checks.append(
                    Check("git repo paths", WARN, f"{len(missing_repos)} not found",
                          "\n".join(f"  {p}" for p in missing_repos[:5]))
                )
            else:
                checks.append(
                    Check("git authors", OK,
                          f"{len(profile.git.authors)} identity(ies), "
                          f"{len(profile.git.repo_paths)} repo(s)")
                )

    sources = models.get("sources.yaml")
    if sources is not None:
        enabled = sources.enabled_aggregators()
        if not enabled:
            checks.append(
                Check("job sources", FAIL, "all disabled",
                      "Enable at least one under aggregators: in config/sources.yaml")
            )
        else:
            ats_n = sources.ats.total_companies()
            checks.append(
                Check("job sources", OK,
                      f"{len(enabled)} aggregator(s): {', '.join(enabled)}"
                      + (f" + {ats_n} ATS board(s)" if ats_n else ""))
            )
        if sources.ats.enabled and sources.ats.total_companies() == 0:
            checks.append(
                Check("ats watchlist", WARN, "empty",
                      "ATS boards are the freshest source — jobs appear there before\n"
                      "aggregators pick them up. Add company slugs under ats: in\n"
                      "config/sources.yaml.")
            )
        if not sources.queries or any(config.contains_placeholder(q) for q in sources.queries):
            checks.append(
                Check("search queries", FAIL, "unset or placeholder",
                      "Set queries: in config/sources.yaml, e.g. ['python backend']")
            )

    voice = models.get("voice.yaml")
    if voice is not None:
        if not voice.pillars:
            checks.append(Check("content pillars", FAIL, "none defined",
                                "Add pillars: to config/voice.yaml"))
        else:
            total_weight = sum(p.weight for p in voice.pillars)
            checks.append(
                Check("content pillars", OK,
                      f"{len(voice.pillars)} pillars, weight sum {total_weight}")
            )
        checks.append(
            Check("banned phrases", OK if voice.banned_phrases else WARN,
                  f"{len(voice.banned_phrases)} phrases, {len(voice.banned_chars)} chars",
                  "An empty ban list lets AI tells through." if not voice.banned_phrases else "")
        )
    return checks


def check_database() -> Check:
    try:
        with db.session() as conn:
            version = conn.execute("PRAGMA user_version").fetchone()[0]
            tables = {
                r["name"]
                for r in conn.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'"
                ).fetchall()
            }
    except (sqlite3.Error, OSError) as exc:
        return Check("database", FAIL, str(exc), f"Check write access to {paths.DATA_DIR}")
    expected = {"jobs", "job_scores", "applications", "activity", "posts", "api_usage", "runs"}
    missing = expected - tables
    if missing:
        return Check("database", FAIL, f"missing tables: {sorted(missing)}",
                     "Delete data/app.db and re-run to rebuild the schema.")
    return Check("database", OK, f"schema v{version}, {len(tables)} tables")


def check_linkedin() -> list[Check]:
    from .linkedin import auth as li_auth

    checks: list[Check] = []
    client_id = os.environ.get("LINKEDIN_CLIENT_ID")
    client_secret = os.environ.get("LINKEDIN_CLIENT_SECRET")
    if not client_id or not client_secret:
        checks.append(
            Check(
                "linkedin app creds",
                FAIL,
                "LINKEDIN_CLIENT_ID / _SECRET unset",
                "One-time setup (~15 min), see README section 'LinkedIn setup':\n"
                "  1. Create a LinkedIn Page (free) — you must be its super admin\n"
                "  2. developer.linkedin.com -> Create app -> associate that Page\n"
                "  3. Settings -> Verify -> open the URL yourself -> approve\n"
                "  4. Products -> add 'Share on LinkedIn' + 'Sign In with LinkedIn\n"
                "     using OpenID Connect' (both self-serve, instant, free)\n"
                "  5. Auth -> redirect URL http://localhost:8765/callback\n"
                "  6. Copy Client ID/Secret into .env (see .env.example)",
            )
        )
    else:
        checks.append(Check("linkedin app creds", OK, "client id + secret present"))

    tok = li_auth.load_tokens()
    if not tok:
        checks.append(
            Check("linkedin token", FAIL, "not authorised",
                  "lgrow linkedin login")
        )
        return checks

    now = datetime.now(timezone.utc)
    access_left = (tok.access_expires_at - now).days if tok.access_expires_at else None
    refresh_left = (tok.refresh_expires_at - now).days if tok.refresh_expires_at else None

    if access_left is not None and access_left < 0:
        status, detail = FAIL, "access token expired"
        hint = "lgrow linkedin refresh"
    elif access_left is not None and access_left < 7:
        status, detail = WARN, f"access token expires in {access_left}d"
        hint = "lgrow linkedin refresh"
    else:
        status = OK
        detail = f"access ok ({access_left}d left)" if access_left is not None else "access ok"
        hint = ""
    if refresh_left is not None and refresh_left < 30:
        status = WARN if status == OK else status
        detail += f", refresh token expires in {refresh_left}d"
        hint = "Re-run `lgrow linkedin login` before the refresh token lapses."
    checks.append(Check("linkedin token", status, detail, hint))

    if tok.person_urn:
        checks.append(Check("linkedin identity", OK, tok.person_urn))
    else:
        checks.append(
            Check("linkedin identity", FAIL, "person URN unknown",
                  "lgrow linkedin login  (needs the openid+profile scopes)")
        )
    return checks


def check_resume() -> Check:
    if not paths.RESUME_BASE.exists():
        return Check(
            "base resume",
            WARN,
            "data/resume_base.docx missing",
            "Tailoring is disabled until you drop your master resume there as\n"
            ".docx (not PDF — we need to edit it).",
        )
    try:
        import docx  # noqa: PLC0415

        doc = docx.Document(str(paths.RESUME_BASE))
        paras = sum(1 for p in doc.paragraphs if p.text.strip())
    except Exception as exc:  # noqa: BLE001
        return Check("base resume", FAIL, f"unreadable: {exc}",
                     "Re-save it as a real .docx from Word/LibreOffice/Google Docs.")
    if paras < 10:
        return Check("base resume", WARN, f"only {paras} non-empty paragraphs",
                     "That looks too short to tailor well — is it the right file?")
    return Check("base resume", OK, f"{paras} paragraphs")


def check_pdf_tooling() -> Check:
    for binary in ("libreoffice", "soffice"):
        if shutil.which(binary):
            return Check("pdf export", OK, f"via {binary}")
    return Check(
        "pdf export",
        WARN,
        "libreoffice not found",
        "Resumes will be produced as .docx only. Install with:\n"
        "    sudo apt install libreoffice-writer",
    )


def check_journal() -> Check:
    if not paths.JOURNAL_PATH.exists():
        return Check("journal", WARN, "no notes yet",
                     "Post drafts need raw material. Add one with:\n"
                     "    lgrow note \"what you figured out today\"")
    lines = [
        ln for ln in paths.JOURNAL_PATH.read_text(encoding="utf-8").splitlines()
        if ln.strip() and not ln.startswith("#")
    ]
    if not lines:
        return Check("journal", WARN, "empty", "lgrow note \"...\"")
    return Check("journal", OK, f"{len(lines)} entries")


# ─── orchestration ───────────────────────────────────────────────────────────


def run(
    probe: bool = False,
    network: bool = False,
    progress: Callable[[str], None] | None = None,
) -> list[Check]:
    """Every check, in order. Results are rendered only after all of them run.

    `progress` exists because the slow checks are silent otherwise. A probe is a
    real model call: up to 90s per attempt, retried with backoff across a
    fallback chain, twice over. Several minutes of a blank terminal is
    indistinguishable from a hang, and the natural response is Ctrl-C — which
    throws away the diagnosis you were waiting for.
    """
    say = progress or (lambda _msg: None)

    checks: list[Check] = [check_python()]
    checks.extend(check_llm_auth())
    checks.extend(check_model_chains())
    checks.extend(check_configs())
    checks.append(check_database())
    checks.append(check_budget())
    checks.extend(check_linkedin())
    checks.append(check_resume())
    checks.append(check_pdf_tooling())
    checks.append(check_journal())

    if probe:
        for tier in ("cheap", "quality"):
            chain = config.load_ai().models.chain(tier)
            model = chain[0] if chain else llm.DEFAULT_MODEL
            say(f"  probing {tier} tier ({model}) — a real call…")
            result = check_gemini_probe(tier)
            say(f"    {result.status} {result.detail}")
            checks.append(result)
    if network:
        say("  checking every job source is reachable…")
        checks.extend(check_sources_reachable())
    return checks


def check_sources_reachable() -> list[Check]:
    from .jobs import sources as job_sources

    checks: list[Check] = []
    for name, probe_fn in job_sources.reachability_probes().items():
        try:
            detail = probe_fn()
        except Exception as exc:  # noqa: BLE001
            checks.append(Check(f"source {name}", WARN, f"unreachable: {exc}",
                                "Transient outages are normal; other sources still run."))
        else:
            checks.append(Check(f"source {name}", OK, detail))
    return checks


def summary(checks: list[Check]) -> tuple[int, int, int]:
    fails = sum(1 for c in checks if c.status == FAIL)
    warns = sum(1 for c in checks if c.status == WARN)
    oks = sum(1 for c in checks if c.status == OK)
    return oks, warns, fails
